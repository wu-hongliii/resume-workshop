import asyncio
import io
import json
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from docx import Document
from docx.shared import RGBColor
from fastapi.testclient import TestClient
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader

from app import ai, demo_package, learning, parsers, providers, renderers, runtime, storage
from app.main import app


def available_cjk_font(size: int) -> ImageFont.FreeTypeFont:
    candidates = (
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    )
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size)
    raise RuntimeError("No supported CJK test font is installed")


def sample_state():
    return {
        "profile": {
            "basics": {
                "name": "Test Candidate",
                "target_role": "AI Engineer",
                "phone": "13800000000",
                "email": "test@example.com",
                "city": "Shanghai",
                "links": "github.com/test",
                "summary": "Builds reliable AI applications.",
                "photo": "",
            },
            "experience": [
                {
                    "company": "Example Co",
                    "role": "AI Engineer",
                    "start": "2024.01",
                    "end": "Present",
                    "location": "Shanghai",
                    "bullets": ["Delivered a knowledge assistant."],
                }
            ],
            "projects": [],
            "education": [],
            "skills": [{"category": "Engineering", "items": ["Python", "RAG"]}],
            "certificates": [],
            "languages": ["Chinese", "English"],
        },
        "resume": {
            "title": "Test Resume",
            "language": "en",
            "template": "ats",
            "page_mode": "one",
            "section_order": ["summary", "experience", "skills", "languages"],
            "jd": "",
            "analysis": None,
            "suggestions": [],
            "questions": [],
        },
    }


class RuntimeTests(unittest.TestCase):
    def test_frozen_data_directories_are_per_user(self):
        with patch.object(runtime.sys, "frozen", True, create=True):
            with patch.object(runtime.sys, "platform", "win32"), patch.dict(
                runtime.os.environ, {"LOCALAPPDATA": r"C:\Users\Friend\AppData\Local"}, clear=False
            ):
                self.assertEqual(
                    runtime._installed_data_dir(),
                    Path(r"C:\Users\Friend\AppData\Local") / "ResumeWorkshop",
                )
            with patch.object(runtime.sys, "platform", "darwin"), patch.object(
                runtime.Path, "home", return_value=Path("/Users/friend")
            ):
                self.assertEqual(
                    runtime._installed_data_dir(),
                    Path("/Users/friend/Library/Application Support/ResumeWorkshop"),
                )

    def test_data_directory_override_wins(self):
        with patch.dict(runtime.os.environ, {"RESUME_WORKSHOP_DATA_DIR": str(Path("build") / "isolated")}, clear=False):
            self.assertEqual(runtime._installed_data_dir(), (Path.cwd() / "build" / "isolated").resolve())


class StorageTests(unittest.TestCase):
    def test_macos_protected_setting_uses_keychain_marker(self):
        original_path = storage.DB_PATH
        with tempfile.TemporaryDirectory() as directory:
            storage.DB_PATH = Path(directory) / "test.db"
            storage.initialize()
            with patch.object(storage.sys, "platform", "darwin"), patch.object(
                storage, "_keychain_set"
            ) as keychain_set, patch.object(storage, "_keychain_get", return_value="sk-from-keychain"):
                storage._set_setting("deepseek_api_key", "sk-secret", protected=True)
                keychain_set.assert_called_once_with("deepseek_api_key", "sk-secret")
                with storage._connect() as db:
                    row = db.execute("SELECT value FROM settings WHERE key = ?", ("deepseek_api_key",)).fetchone()
                self.assertEqual(row["value"], storage.KEYCHAIN_MARKER)
                self.assertEqual(storage._get_setting("deepseek_api_key", protected=True), "sk-from-keychain")
        storage.DB_PATH = original_path

    def test_state_versions_backup_and_key(self):
        original_path = storage.DB_PATH
        with tempfile.TemporaryDirectory() as directory:
            storage.DB_PATH = Path(directory) / "test.db"
            storage.initialize()
            payload = sample_state()
            storage.save_state(payload, "first version")
            self.assertEqual(storage.get_state()["profile"]["basics"]["name"], "Test Candidate")
            self.assertEqual(len(storage.list_versions()), 1)
            storage.set_api_key("sk-test-secret")
            self.assertEqual(storage.get_api_key(), "sk-test-secret")
            storage.set_search_settings("brave", "search-test-secret")
            search = storage.get_search_settings()
            self.assertEqual(search["provider"], "brave")
            self.assertEqual(search["api_key"], "search-test-secret")
            storage.set_learning_resource_cache("bridge", [{"title": "桥梁工程"}])
            self.assertEqual(storage.get_learning_resource_cache("bridge")[0]["title"], "桥梁工程")
            storage.set_learning_resource_feedback("https://www.bilibili.com/video/BV1234567890/", "useful")
            self.assertEqual(storage.get_learning_resource_feedback("https://www.bilibili.com/video/BV1234567890/"), "useful")
            backup = storage.export_backup()
            self.assertNotIn("sk-test-secret", json.dumps(backup))
            self.assertNotIn("search-test-secret", json.dumps(backup))
            storage.import_backup(backup)
            self.assertEqual(storage.get_state()["resume"]["title"], "Test Resume")
            self.assertEqual(len(backup["learning_feedback"]), 1)
        storage.DB_PATH = original_path

    def test_ai_provider_settings_migrate_and_keep_provider_keys_separate(self):
        original_path = storage.DB_PATH
        with tempfile.TemporaryDirectory() as directory:
            storage.DB_PATH = Path(directory) / "test.db"
            storage.initialize()
            storage._set_setting("deepseek_api_key", "legacy-deepseek", protected=True)
            self.assertEqual(storage.get_ai_settings()["api_key"], "legacy-deepseek")
            storage.set_ai_settings("openai", "https://api.openai.com/v1", "gpt-5.2", "openai-secret")
            self.assertEqual(storage.get_ai_settings()["provider"], "openai")
            self.assertEqual(storage.get_ai_settings()["api_key"], "openai-secret")
            self.assertEqual(storage.get_ai_settings("deepseek")["api_key"], "legacy-deepseek")
            backup_text = json.dumps(storage.export_backup())
            self.assertNotIn("legacy-deepseek", backup_text)
            self.assertNotIn("openai-secret", backup_text)
        storage.DB_PATH = original_path


class ProviderTests(unittest.TestCase):
    def test_provider_catalog_and_validation(self):
        ids = {item["id"] for item in providers.public_provider_options()}
        self.assertEqual(
            ids,
            {"deepseek", "openai", "anthropic", "minimax", "glm", "kimi", "mimo", "ollama", "lmstudio", "custom"},
        )
        local = providers.normalize_config({"provider": "ollama"})
        self.assertEqual(local["base_url"], "http://127.0.0.1:11434/v1")
        with self.assertRaisesRegex(providers.ProviderError, "API Key"):
            providers.normalize_config({"provider": "openai"})
        with self.assertRaisesRegex(providers.ProviderError, "不能包含账号或密码"):
            providers.normalize_config({
                "provider": "custom", "base_url": "https://user:pass@example.com/v1", "model": "test", "api_key": "key"
            })
        with self.assertRaisesRegex(providers.ProviderError, "必须使用 HTTPS"):
            providers.normalize_config({
                "provider": "custom", "base_url": "http://models.example.com/v1", "model": "test", "api_key": "key"
            })
        local_custom = providers.normalize_config({
            "provider": "custom", "base_url": "http://127.0.0.1:9000/v1", "model": "test", "api_key": "local-placeholder"
        })
        self.assertEqual(local_custom["base_url"], "http://127.0.0.1:9000/v1")

    def test_anthropic_response_is_normalized(self):
        class FakeResponse:
            status_code = 200

            @staticmethod
            def json():
                return {
                    "content": [{"type": "text", "text": '{"ok":true}'}],
                    "stop_reason": "end_turn",
                    "model": "claude-sonnet-5",
                }

        class FakeClient:
            def __init__(self, **_kwargs):
                self.request = None

            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                return None

            async def post(self, url, headers, json):
                self.request = (url, headers, json)
                self.__class__.last_request = self.request
                return FakeResponse()

        config = {
            "provider": "anthropic", "base_url": "https://api.anthropic.com", "model": "claude-sonnet-5", "api_key": "secret"
        }
        request = {
            "max_tokens": 100,
            "messages": [{"role": "system", "content": "system"}, {"role": "user", "content": "user"}],
        }
        with patch("app.providers.httpx.AsyncClient", FakeClient):
            result = asyncio.run(providers.request_completion(config, request))
        url, headers, body = FakeClient.last_request
        self.assertEqual(url, "https://api.anthropic.com/v1/messages")
        self.assertEqual(headers["x-api-key"], "secret")
        self.assertEqual(body["system"], "system")
        self.assertEqual(result["choices"][0]["message"]["content"], '{"ok":true}')


class StorageRestoreTests(unittest.TestCase):
    def test_restore_keeps_real_pre_restore_snapshot(self):
        original_path = storage.DB_PATH
        with tempfile.TemporaryDirectory() as directory:
            storage.DB_PATH = Path(directory) / "test.db"
            storage.initialize()
            old = sample_state()
            old["resume"]["title"] = "Old version"
            storage.save_state(old, "old")
            old_id = storage.list_versions()[0]["id"]
            current = sample_state()
            current["resume"]["title"] = "Current version"
            storage.save_state(current, "current")

            restored = storage.restore_version(old_id)
            self.assertEqual(restored["resume"]["title"], "Old version")
            snapshot = next(item for item in storage._all_versions() if item["label"] == "恢复历史版本前的自动快照")
            self.assertEqual(json.loads(snapshot["payload"])["resume"]["title"], "Current version")
        storage.DB_PATH = original_path

    def test_backup_import_deduplicates_and_versions_are_bounded(self):
        original_path = storage.DB_PATH
        with tempfile.TemporaryDirectory() as directory:
            storage.DB_PATH = Path(directory) / "test.db"
            storage.initialize()
            for index in range(storage.MAX_VERSIONS + 5):
                payload = sample_state()
                payload["resume"]["title"] = f"Version {index}"
                storage.save_state(payload, f"version {index}")
            self.assertEqual(len(storage._all_versions()), storage.MAX_VERSIONS)
            backup = storage.export_backup()
            storage.import_backup(backup)
            count_after_first = len(storage._all_versions())
            storage.import_backup(backup)
            self.assertEqual(len(storage._all_versions()), count_after_first)
        storage.DB_PATH = original_path


class ParserRendererTests(unittest.TestCase):
    def test_office_zip_bomb_is_rejected_before_parsing(self):
        buffer = io.BytesIO()
        with zipfile.ZipFile(buffer, "w", compression=zipfile.ZIP_DEFLATED) as archive:
            archive.writestr("word/document.xml", b"A" * (31 * 1024 * 1024))
        with self.assertRaisesRegex(ValueError, "单个内容过大"):
            parsers.parse_docx(buffer.getvalue())

    def test_oversized_image_dimensions_are_rejected(self):
        image = Image.new("RGB", (12_001, 1), "white")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        with self.assertRaisesRegex(ValueError, "图片尺寸过大"):
            parsers.parse_image(buffer.getvalue())

    def test_docx_roundtrip(self):
        content = renderers.render_docx(sample_state())
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            self.assertIn("word/document.xml", archive.namelist())
        text = parsers.parse_docx(content)
        self.assertIn("Test Candidate", text)
        self.assertIn("AI Engineer", text)

    def test_html_and_pdf(self):
        markup = renderers.render_resume_html(sample_state())
        self.assertIn("Test Candidate", markup)
        self.assertIn("template-ats", markup)
        self.assertIn("resume-fit", markup)
        content = asyncio.run(renderers.render_pdf(sample_state()))
        self.assertTrue(content.startswith(b"%PDF"))

    def test_one_page_pdf_fits_and_auto_mode_paginates(self):
        payload = sample_state()
        base = payload["profile"]["experience"][0]
        payload["profile"]["experience"] = [
            {**base, "company": f"Example Co {index}", "bullets": ["Delivered a detailed workflow and documented validation results." for _ in range(4)]}
            for index in range(18)
        ]
        payload["resume"]["page_mode"] = "one"
        with self.assertRaises(renderers.ResumeLayoutError):
            asyncio.run(renderers.render_pdf(payload))
        payload["resume"]["page_mode"] = "auto"
        multi_page = asyncio.run(renderers.render_pdf(payload))
        self.assertGreater(len(PdfReader(io.BytesIO(multi_page)).pages), 1)

    def test_tailored_content_keeps_base_identity(self):
        payload = sample_state()
        payload["resume"]["tailored_profile"] = {
            **payload["profile"],
            "basics": {
                **payload["profile"]["basics"],
                "name": "Wrong Name",
                "phone": "000",
                "summary": "Targeted AI delivery summary.",
            },
        }
        effective = renderers.effective_profile(payload)
        self.assertEqual(effective["basics"]["name"], "Test Candidate")
        self.assertEqual(effective["basics"]["phone"], "13800000000")
        self.assertEqual(effective["basics"]["summary"], "Targeted AI delivery summary.")
        self.assertIn("Targeted AI delivery summary.", renderers.render_resume_html(payload))

    def test_chinese_job_screenshot_ocr(self):
        lines = [
            "\u804c\u4f4d\u63cf\u8ff0",
            "\u8d1f\u8d23\u5927\u6a21\u578b\u5e94\u7528\u7684\u9700\u6c42\u5206\u6790\u3001\u65b9\u6848\u8bbe\u8ba1\u4e0e\u9879\u76ee\u4ea4\u4ed8",
            "\u642d\u5efa RAG \u77e5\u8bc6\u5e93\uff0c\u4f18\u5316\u68c0\u7d22\u51c6\u786e\u7387\u548c\u7528\u6237\u4f53\u9a8c",
            "\u4efb\u804c\u8981\u6c42",
        ]
        image = Image.new("RGB", (1400, 520), "white")
        draw = ImageDraw.Draw(image)
        font = available_cjk_font(30)
        for index, line in enumerate(lines):
            draw.text((60, 40 + index * 105), line, font=font, fill=(75, 75, 75))
        buffer = io.BytesIO()
        image.save(buffer, "PNG")
        text = parsers.parse_image(buffer.getvalue())
        self.assertIn(lines[0], text)
        self.assertIn("RAG", text)
        self.assertIn(lines[3], text)
        parsed = parsers.parse_file("jd.png", buffer.getvalue())
        self.assertIsInstance(parsed["ocr_confidence"], float)
        self.assertGreater(parsed["ocr_confidence"], 0)


class PrivacyTests(unittest.TestCase):
    def test_redact_and_restore(self):
        source = "Call 13800000000 or test@example.com"
        redacted, replacements = ai.redact_sensitive(source)
        self.assertNotIn("13800000000", redacted)
        self.assertNotIn("test@example.com", redacted)
        self.assertEqual(ai.restore_sensitive(redacted, replacements), source)

    def test_profile_for_ai_excludes_direct_identity_and_photo(self):
        profile = sample_state()["profile"]
        profile["basics"]["photo"] = "data:image/png;base64,very-private-photo"
        safe = ai.profile_for_ai(profile)
        serialized = json.dumps(safe, ensure_ascii=False)
        for private_value in ("Test Candidate", "13800000000", "test@example.com", "Shanghai", "github.com/test", "very-private-photo"):
            self.assertNotIn(private_value, serialized)
        self.assertIn("Builds reliable AI applications", serialized)
        self.assertIn("Example Co", serialized)

    def test_match_score_is_calculated_from_evidence_weights(self):
        async def fake_chat(*_args, **_kwargs):
            return {
                "summary": "Evidence based",
                "requirements": [
                    {"requirement": "Python", "category": "must", "status": "met", "evidence": "Project"},
                    {"requirement": "Trading", "category": "responsibility", "status": "partial", "evidence": "Practice"},
                    {"requirement": "CFA", "category": "preferred", "status": "missing", "evidence": "None"},
                ],
            }
        with patch("app.ai._chat", side_effect=fake_chat):
            result = asyncio.run(ai.analyze_jd("key", sample_state()["profile"], "JD"))
        self.assertEqual(result["score"], 67)
        self.assertIn("3 条岗位要求", result["score_basis"])

    def test_targeted_generation_protects_identity(self):
        async def fake_chat(*_args, **_kwargs):
            return {
                "target_role": "FDE",
                "summary": "Targeted",
                "experience": [{"source_index": 0, "bullets": ["Delivered an AI application"]}],
                "projects": [],
                "skills_order": [0],
                "generated_projects": [{"name": "AI Delivery Lab", "role": "个人项目", "start": "", "end": "", "technologies": "Python", "bullets": ["Built a delivery workflow"]}],
                "generated_skills": [{"category": "Target skills", "items": ["Python", "RAG"]}],
            }

        with patch("app.ai._chat", side_effect=fake_chat):
            result = asyncio.run(ai.generate_targeted_profile("key", sample_state()["profile"], "AI Engineer", "Customer delivery"))
        generated = result["profile"]
        self.assertEqual(generated["basics"]["name"], "Test Candidate")
        self.assertEqual(generated["basics"]["phone"], "13800000000")
        self.assertEqual(generated["basics"]["summary"], "Targeted")
        self.assertEqual(generated["experience"][0]["company"], "Example Co")
        self.assertEqual(generated["experience"][0]["bullets"], ["Delivered an AI application"])
        self.assertEqual(generated["projects"][0]["role"], "个人项目")
        self.assertEqual(generated["skills"][0]["category"], "Target skills")

    def test_compact_plan_explicit_selection_omits_unselected_sources(self):
        profile = sample_state()["profile"]
        plan = {
            "summary": "Role-focused summary",
            "experience": [{"source_index": 99, "bullets": ["Invented role"]}],
            "skills_order": [0, 99],
            "additional_projects": [],
        }
        result = ai.assemble_targeted_profile(profile, plan)
        self.assertEqual(result["experience"], [])
        self.assertEqual(result["projects"], [])
        self.assertEqual(result["education"], profile["education"])
        self.assertEqual(result["basics"]["name"], "Test Candidate")

    def test_compact_plan_missing_section_preserves_base_as_safe_fallback(self):
        profile = sample_state()["profile"]
        result = ai.assemble_targeted_profile(profile, {"summary": "Role-focused summary"})
        self.assertEqual(result["experience"], profile["experience"])
        self.assertEqual(result["projects"], profile["projects"])

    def test_refinement_plan_applies_answer_without_changing_identity(self):
        profile = sample_state()["profile"]
        result = ai.apply_refinement_plan(profile, {"operations": [
            {"op": "replace_summary", "value": "Refined summary"},
            {"op": "append_experience_bullet", "item_index": 0, "value": "Added from user answer"},
            {"op": "add_skill", "payload": {"category": "Trading", "items": ["Risk control"]}},
        ]})
        self.assertEqual(result["basics"]["name"], "Test Candidate")
        self.assertEqual(result["basics"]["summary"], "Refined summary")
        self.assertIn("Added from user answer", result["experience"][0]["bullets"])
        self.assertEqual(result["skills"][-1]["category"], "Trading")

    def test_refinement_can_remove_denied_project_and_reports_change(self):
        profile = sample_state()["profile"]
        profile["projects"] = [{"name": "Untrue simulation", "role": "模拟项目", "bullets": ["Backtested a strategy"]}]
        result = ai.apply_refinement_plan(profile, {"operations": [{"op": "remove_project", "item_index": 0}]})
        changes = ai.describe_profile_changes(profile, result)
        self.assertEqual(result["projects"], [])
        self.assertEqual(changes, [{"path": "projects", "label": "项目经历数量（1→0）"}])

    def test_refine_with_answer_returns_verified_change_list(self):
        profile = sample_state()["profile"]
        profile["projects"] = [{"name": "Untrue simulation", "role": "模拟项目", "bullets": ["Backtested a strategy"]}]

        async def fake_chat(*_args, **_kwargs):
            return {"operations": [{"op": "remove_project", "item_index": 0}]}

        with patch("app.ai._chat", side_effect=fake_chat):
            response = asyncio.run(ai.refine_with_answer("key", profile, "Trader JD", "Did you backtest?", "No", []))
        self.assertEqual(response["profile"]["projects"], [])
        self.assertEqual(response["changes"][0]["path"], "projects")

    def test_add_resume_content_appends_to_relevant_project_without_touching_jobs(self):
        profile = sample_state()["profile"]
        profile["projects"] = [{"name": "交易复盘工具", "role": "个人项目", "bullets": ["记录每日交易"]}]

        async def fake_chat(*_args, **_kwargs):
            return {"operations": [
                {"op": "append_project_bullet", "item_index": 0, "value": "使用通达信制作个股与板块 RPS 指标，辅助复盘选股"},
                {"op": "append_experience_bullet", "item_index": 0, "value": "不应加入工作经历"},
                {"op": "add_skill", "payload": {"category": "交易工具", "items": ["通达信", "RPS 指标"]}},
            ]}

        with patch("app.ai._chat", side_effect=fake_chat):
            response = asyncio.run(ai.add_resume_content("key", profile, "交易员 JD", "我制作了 RPS 指标"))
        self.assertIn("通达信", response["profile"]["projects"][0]["bullets"][-1])
        self.assertNotIn("不应加入工作经历", response["profile"]["experience"][0]["bullets"])
        self.assertEqual(response["profile"]["skills"][-1]["category"], "交易工具")
        self.assertEqual(response["profile"]["basics"]["name"], "Test Candidate")

    def test_structural_edit_can_remove_and_regenerate_project(self):
        profile = sample_state()["profile"]
        profile["projects"] = [
            {"name": "模拟交易练习", "role": "个人项目", "start": "", "end": "", "technologies": "模拟盘", "bullets": ["测试规则"]},
            {"name": "保留项目", "role": "个人项目", "start": "", "end": "", "technologies": "Python", "bullets": ["保留内容"]},
        ]

        async def fake_chat(*_args, **_kwargs):
            return {"operations": [
                {"op": "replace_project", "item_index": 0, "payload": {
                    "name": "通达信指标与复盘工具", "role": "个人项目", "technologies": "通达信",
                    "bullets": ["制作个股与板块 RPS 三线红指标", "使用月线反转指标辅助复盘选股"],
                }},
                {"op": "remove_skill", "item_index": 0},
                {"op": "add_experience", "payload": {"company": "虚构公司"}},
            ]}

        with patch("app.ai._chat", side_effect=fake_chat):
            response = asyncio.run(ai.add_resume_content("key", profile, "交易员 JD", "删掉模拟项目并重新生成通达信项目", "adjust"))
        result = response["profile"]
        self.assertEqual(result["projects"][0]["name"], "通达信指标与复盘工具")
        self.assertEqual(result["projects"][1]["name"], "保留项目")
        self.assertEqual(result["experience"], profile["experience"])
        self.assertEqual(result["skills"], [])

    def test_edit_selected_text_returns_only_replacement(self):
        async def fake_chat(*_args, **_kwargs):
            return {"replacement_text": "面向企业客户完成 AI 应用交付"}

        with patch("app.ai._chat", side_effect=fake_chat) as chat:
            response = asyncio.run(ai.edit_selected_text(
                "key", sample_state()["profile"], "FDE JD", "basics.summary", "个人概述",
                "具备应用交付经验", "应用交付", 2, 6, "改得更贴合 FDE",
            ))
        self.assertEqual(response, {"replacement_text": "面向企业客户完成 AI 应用交付"})
        self.assertIn("只改写用户选中的片段", chat.call_args.args[1])
        self.assertIn("应用交付", chat.call_args.args[2])

    def test_edit_selected_text_rejects_whole_field_repeated_into_selection(self):
        full_text = "自律上进，勤于思考，执行力强。2025年2月起全职交易，实现171.27%收益，最大回撤-5.90%。"
        selected = "171.27%收益"
        start = full_text.index(selected)

        async def fake_chat(*_args, **_kwargs):
            return {"replacement_text": full_text}

        with patch("app.ai._chat", side_effect=fake_chat):
            with self.assertRaisesRegex(ai.AIError, "整段或重复内容"):
                asyncio.run(ai.edit_selected_text(
                    "key", sample_state()["profile"], "交易员 JD", "basics.summary", "个人概述",
                    full_text, selected, start, start + len(selected), "补充收益区间",
                ))

    def test_edit_whole_field_rejects_new_repeated_paragraph(self):
        original = "具备交易复盘与风险控制经验。"
        repeated = "持续跟踪市场并复盘交易决策，能够保持纪律执行。持续跟踪市场并复盘交易决策，能够保持纪律执行。"

        async def fake_chat(*_args, **_kwargs):
            return {"replacement_text": repeated}

        with patch("app.ai._chat", side_effect=fake_chat):
            with self.assertRaisesRegex(ai.AIError, "重复内容"):
                asyncio.run(ai.edit_selected_text(
                    "key", sample_state()["profile"], "交易员 JD", "basics.summary", "个人概述",
                    original, "", 0, 0, "重写个人概述",
                ))

    def test_json_parser_accepts_wrapped_object(self):
        value = ai._extract_json("Result follows:\n{\"ok\": true}\nDone")
        self.assertEqual(value, {"ok": True})

    def test_chat_retries_empty_json_response(self):
        responses = [
            {"choices": [{"finish_reason": "stop", "message": {"content": ""}}]},
            {"choices": [{"finish_reason": "stop", "message": {"content": '{"ok": true}'}}]},
        ]
        with patch("app.ai._request_completion", side_effect=responses) as request:
            result = asyncio.run(ai._chat("key", "Return JSON", "input", max_tokens=1000))
        self.assertEqual(result, {"ok": True})
        self.assertEqual(request.call_count, 2)
        self.assertEqual(request.call_args_list[1].args[1]["max_tokens"], 2000)


class ApiTests(unittest.TestCase):
    def test_update_check_is_explicit_and_returns_latest_release(self):
        class FakeResponse:
            status_code = 200

            @staticmethod
            def json():
                return {
                    "tag_name": "v0.3.1",
                    "html_url": "https://github.com/wu-hongliii/resume-workshop/releases/tag/v0.3.1",
                }

        class FakeClient:
            def __init__(self, **_kwargs):
                pass

            async def __aenter__(self):
                return self

            async def __aexit__(self, *_args):
                return None

            async def get(self, url, headers):
                self.test_url = url
                self.test_headers = headers
                return FakeResponse()

        with patch("app.main.httpx.AsyncClient", FakeClient):
            with TestClient(app) as client:
                response = client.get("/api/update-check")
        self.assertEqual(response.status_code, 200)
        self.assertTrue(response.json()["update_available"])
        self.assertEqual(response.json()["latest"], "0.3.1")

    def test_demo_package_contains_two_complete_fictional_workflows(self):
        package = demo_package.build_demo_package()
        with zipfile.ZipFile(io.BytesIO(package)) as archive:
            names = archive.namelist()
            expected_suffixes = (
                "00-开始前请读.txt",
                "01-旧简历-林知远-AI应用工程师-演示专用.docx",
                "02-目标岗位JD-FDE-演示专用.txt",
                "01-旧简历-周桥-桥涵设计工程师-演示专用.docx",
                "03-目标岗位JD截图-OCR演示专用.png",
                "编辑部单栏-演示模板.docx",
            )
            for suffix in expected_suffixes:
                self.assertTrue(any(name.endswith(suffix) for name in names), suffix)

            old_resume_name = next(name for name in names if name.endswith("01-旧简历-林知远-AI应用工程师-演示专用.docx"))
            parsed_resume = parsers.parse_docx(archive.read(old_resume_name))
            self.assertIn("林知远", parsed_resume)
            self.assertIn("演示专用", parsed_resume)

            template_name = next(name for name in names if name.endswith("编辑部单栏-演示模板.docx"))
            template = parsers.analyze_template(template_name, archive.read(template_name))["template_spec"]
            self.assertEqual(template["source_type"], "docx")
            self.assertIn("编辑部单栏", template["name"])

            screenshot_name = next(name for name in names if name.endswith("03-目标岗位JD截图-OCR演示专用.png"))
            with Image.open(io.BytesIO(archive.read(screenshot_name))) as screenshot:
                self.assertEqual(screenshot.format, "PNG")
                self.assertGreater(screenshot.width, 1000)

    def test_demo_package_endpoint_downloads_zip_without_state_write(self):
        with patch("app.main.storage.save_state") as save_state:
            with TestClient(app) as client:
                response = client.get("/api/demo-package")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.headers["content-type"], "application/zip")
        self.assertIn("filename*=UTF-8", response.headers["content-disposition"])
        self.assertEqual(zipfile.ZipFile(io.BytesIO(response.content)).testzip(), None)
        save_state.assert_not_called()

    def test_health_preview_and_missing_key(self):
        with TestClient(app) as client:
            health = client.get("/api/health")
            self.assertTrue(health.json()["ok"])
            self.assertIn("frame-ancestors 'none'", health.headers.get("content-security-policy", ""))
            home = client.get("/")
            self.assertIn("no-store", home.headers.get("cache-control", ""))
            self.assertEqual(home.headers.get("x-content-type-options"), "nosniff")
            preview = client.post("/api/preview", json=sample_state())
            self.assertEqual(preview.status_code, 200)
            self.assertIn("Test Candidate", preview.text)
            with patch("app.main.storage.get_api_key", return_value=""):
                response = client.post("/api/ai/analyze", json={"profile": {}, "jd": "Python"})
            self.assertEqual(response.status_code, 400)

    def test_local_boundary_and_state_shape_validation(self):
        with TestClient(app) as client:
            blocked = client.put(
                "/api/state",
                headers={"Origin": "https://malicious.example"},
                json={"payload": sample_state()},
            )
            self.assertEqual(blocked.status_code, 403)
            malformed = client.put("/api/state", json={"payload": {"profile": {}}})
            self.assertEqual(malformed.status_code, 422)

    def test_custom_template_is_rebuilt_without_importing_sample_content(self):
        document = Document()
        run = document.add_paragraph("示例人物不应进入档案").add_run("模板标题")
        run.font.name = "Microsoft YaHei"
        run.font.color.rgb = RGBColor(139, 74, 50)
        buffer = io.BytesIO(); document.save(buffer)
        with TestClient(app) as client:
            response = client.post(
                "/api/template/analyze?filename=sample.docx",
                content=buffer.getvalue(),
                headers={"Content-Type": "application/octet-stream"},
            )
        self.assertEqual(response.status_code, 200)
        spec = response.json()["template_spec"]
        self.assertEqual(spec["accent"], "#8b4a32")
        self.assertNotIn("content", spec)
        self.assertEqual(spec["name"], "sample")

    def test_folder_import_supports_text_and_xlsx(self):
        text = parsers.parse_file("project/readme.md", "项目复盘".encode())
        self.assertEqual(text["text"], "项目复盘")
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "evidence.xlsx"
            document = zipfile.ZipFile(path, "w")
            document.writestr("xl/worksheets/sheet1.xml", '<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"><sheetData><row><c t="inlineStr"><is><t>研究结论</t></is></c></row></sheetData></worksheet>')
            document.close()
            parsed = parsers.parse_file(path.name, path.read_bytes())
        self.assertIn("研究结论", parsed["text"])

    def test_learning_path_adds_direct_course_links(self):
        async def fake_path(*_args, **_kwargs):
            return {"summary": "先补基础", "stages": [{"title": "Python", "keywords": ["Python 入门"]}]}

        async def fake_resources(result, *_args):
            result["stages"][0]["resources"] = [{
                "title": "Python 零基础教程", "platform": "哔哩哔哩",
                "url": "https://www.bilibili.com/video/BV1R7411F7JV/", "kind": "直达课程",
            }]
            return result

        async def fake_rank(_key, result):
            return result

        with patch("app.main.storage.get_api_key", return_value="key"), patch("app.main.storage.get_search_settings", return_value={"provider": "", "api_key": ""}), patch("app.main.ai.generate_learning_path", side_effect=fake_path), patch("app.main.learning.attach_direct_resources", side_effect=fake_resources), patch("app.main.ai.rank_learning_resources", side_effect=fake_rank):
            with TestClient(app) as client:
                response = client.post("/api/ai/learning-path", json={"profile": {}, "jd": "Python", "analysis": {}})
        self.assertEqual(response.status_code, 200)
        resources = response.json()["stages"][0]["resources"]
        self.assertEqual(len(resources), 1)
        self.assertIn("/video/", resources[0]["url"])
        self.assertNotIn("search", resources[0]["url"])

    def test_learning_resource_filter_rejects_search_pages(self):
        direct = learning._direct_resource(
            "LangChain Agent 教程", "https://www.bilibili.com/video/BV1rv7A6oEeP/"
        )
        search = learning._direct_resource(
            "搜索", "https://search.bilibili.com/all?keyword=LangChain"
        )
        self.assertEqual(direct["platform"], "哔哩哔哩")
        self.assertIsNone(search)

    def test_learning_resource_supports_smartedu_course_pages(self):
        resource = learning._direct_resource(
            "桥梁工程 I", "https://higher.smartedu.cn/course/69a20dea95df98bb276c9fe6"
        )
        self.assertEqual(resource["platform"], "国家高等教育智慧教育平台")
        self.assertEqual(resource["url"], "https://higher.smartedu.cn/course/69a20dea95df98bb276c9fe6")

    def test_learning_resources_are_interleaved_across_platforms(self):
        resources = [
            {"platform": "哔哩哔哩", "url": "b1"},
            {"platform": "哔哩哔哩", "url": "b2"},
            {"platform": "中国大学 MOOC", "url": "m1"},
            {"platform": "国家高等教育智慧教育平台", "url": "s1"},
        ]
        diversified = learning.diversify_resources(resources)
        self.assertEqual([item["url"] for item in diversified[:3]], ["b1", "m1", "s1"])

    def test_learning_resource_pool_starts_with_diverse_platforms(self):
        result = {"stages": [{"resource_pool": [
            {"platform": "哔哩哔哩", "url": "b1"},
            {"platform": "哔哩哔哩", "url": "b2"},
            {"platform": "中国大学 MOOC", "url": "m1"},
            {"platform": "国家高等教育智慧教育平台", "url": "s1"},
        ]}]}
        learning.finalize_resource_pools(result)
        self.assertEqual([item["url"] for item in result["stages"][0]["resources"]], ["b1", "m1", "s1"])

    def test_ai_ranking_keeps_unselected_candidates_for_refresh(self):
        result = {"stages": [{"title": "Python", "resource_candidates": [
            {"candidate_id": "S1R1", "title": "B站课程", "platform": "哔哩哔哩", "url": "https://www.bilibili.com/video/BV1234567890/"},
            {"candidate_id": "S1R2", "title": "大学课程", "platform": "中国大学 MOOC", "url": "https://www.icourse163.org/course/detail.htm?cid=268001"},
            {"candidate_id": "S1R3", "title": "国家平台课程", "platform": "国家高等教育智慧教育平台", "url": "https://higher.smartedu.cn/course/68c0987da9f4619f8f00c004"},
            {"candidate_id": "S1R4", "title": "慕课课程", "platform": "慕课网", "url": "https://www.imooc.com/learn/177"},
        ]}]}

        async def fake_chat(*_args, **_kwargs):
            return {"stages": [{"stage_index": 0, "selections": [{"candidate_id": "S1R1", "reason": "匹配"}]}]}

        with patch("app.ai._chat", side_effect=fake_chat):
            ranked = asyncio.run(ai.rank_learning_resources("key", result))
        self.assertEqual(len(ranked["stages"][0]["resource_pool"]), 4)
        self.assertEqual({item["platform"] for item in ranked["stages"][0]["resource_pool"]}, {"哔哩哔哩", "中国大学 MOOC", "国家高等教育智慧教育平台", "慕课网"})

    def test_learning_resources_match_stage_instead_of_returning_platform_searches(self):
        resources = learning.find_direct_resources({
            "title": "AI 与 Agent 开发入门",
            "goal": "理解 LangChain、RAG 和智能体开发",
            "keywords": ["LangChain", "RAG", "Agent"],
        })
        self.assertGreaterEqual(len(resources), 2)
        self.assertTrue(all(learning.is_supported_direct_url(item["url"]) for item in resources))
        self.assertTrue(all("search" not in item["url"] for item in resources))
        self.assertGreaterEqual(len({item["platform"] for item in resources}), 2)

    def test_material_batches_keep_long_documents_beyond_old_limit(self):
        files = [
            {"filename": f"part-{index}.docx", "text": str(index) * 20_000}
            for index in range(4)
        ]
        batches = ai.build_material_batches(files)
        self.assertGreaterEqual(len(batches), 3)
        self.assertGreater(sum(len(item["text"]) for batch in batches for item in batch), 45_000)

    def test_general_material_case_uses_engineering_labels_and_evidence(self):
        responses = iter([
            {"facts": [{"fact": "完成箱涵施工图", "evidence_filename": "设计总结.docx", "evidence_quote": "完成箱涵施工图设计", "confidence": 0.95}], "material_signals": []},
            {"material_type": "engineering", "classification_confidence": 0.91, "alternatives": [], "name": "道路箱涵设计", "role": "设计人员", "summary": "完成道路箱涵设计", "methods": "规范校核、结构计算", "structure": "方案比选至施工图", "highlights": ["完成结构验算"], "bullets": ["完成箱涵施工图设计"], "evidence_items": [{"statement": "完成箱涵施工图", "filename": "设计总结.docx", "quote": "完成箱涵施工图设计", "confidence": 0.95}, {"statement": "虚构依据", "filename": "设计总结.docx", "quote": "原文不存在的内容", "confidence": 0.2}], "uncertainties": []},
        ])

        async def fake_chat(*_args, **_kwargs):
            return next(responses)

        with patch("app.ai._chat", side_effect=fake_chat):
            draft = asyncio.run(ai.draft_project_case("key", {}, "市政桥涵设计", "材料", [{"filename": "设计总结.docx", "text": "完成箱涵施工图设计"}]))
        self.assertEqual(draft["material_type"], "engineering")
        self.assertEqual(draft["labels"]["methods"], "规范与工具")
        self.assertEqual(draft["evidence_items"][0]["filename"], "设计总结.docx")
        self.assertEqual(len(draft["evidence_items"]), 1)

    def test_remote_course_search_keeps_only_specific_course_pages(self):
        original_path = storage.DB_PATH
        with tempfile.TemporaryDirectory() as directory:
            storage.DB_PATH = Path(directory) / "test.db"

            async def fake_brave(*_args):
                return [
                    {"title": "市政桥涵设计教程", "url": "https://www.bilibili.com/video/BV1234567890/", "description": "箱涵设计与结构计算"},
                    {"title": "搜索结果", "url": "https://search.bilibili.com/all?keyword=桥涵"},
                ]

            result = {"stages": [{"title": "桥涵设计基础", "goal": "掌握箱涵设计", "keywords": ["市政桥涵设计"]}]}
            with patch("app.learning._search_brave", side_effect=fake_brave):
                attached = asyncio.run(learning.attach_direct_resources(result, "brave", "key"))
            resources = attached["stages"][0]["resources"]
            self.assertTrue(all(learning.is_supported_direct_url(item["url"]) for item in resources))
            self.assertTrue(any("/video/" in item["url"] for item in resources))
            self.assertTrue(all("search" not in item["url"] for item in resources))
        storage.DB_PATH = original_path


if __name__ == "__main__":
    unittest.main()
