import os
import re
from pathlib import Path

from docx import Document
from playwright.sync_api import expect, sync_playwright


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "exports" / "e2e"
OUTPUT.mkdir(parents=True, exist_ok=True)


def isolated_test_url() -> str:
    url = os.environ.get("RESUME_TEST_URL", "").strip()
    writes_allowed = os.environ.get("RESUME_E2E_ALLOW_STATE_WRITES") == "1"
    if not url or not writes_allowed:
        raise RuntimeError(
            "e2e_smoke 会写入应用状态。请先启动使用隔离数据库的测试服务，"
            "再同时设置 RESUME_TEST_URL 和 RESUME_E2E_ALLOW_STATE_WRITES=1；"
            "不要指向保存真实简历的默认服务。"
        )
    return url


def main() -> None:
    base_url = isolated_test_url()
    console_errors: list[str] = []
    fixture = OUTPUT / "import-fixture.docx"
    second_fixture = OUTPUT / "import-fixture-2.docx"
    document = Document()
    document.add_paragraph("Resume Import Smoke Test")
    document.add_paragraph("AI Engineer")
    document.save(fixture)
    second_document = Document()
    second_document.add_paragraph("Second Resume Evidence File")
    second_document.add_paragraph("Trading Research")
    second_document.save(second_fixture)
    with sync_playwright() as playwright:
        browser = playwright.chromium.launch(headless=True)
        page = browser.new_page(viewport={"width": 1440, "height": 1000}, device_scale_factor=1)
        page.add_init_script("localStorage.setItem('resume-workshop:onboarding:v2', 'done'); localStorage.setItem('resume-workshop:terms:v1', 'accepted'); localStorage.setItem('resume-workshop:ai-consent:deepseek:https://api.deepseek.com:v1', 'accepted')")
        page.on("console", lambda message: console_errors.append(message.text) if message.type == "error" else None)
        page.on("pageerror", lambda error: console_errors.append(str(error)))

        page.goto(base_url)
        page.wait_for_load_state("networkidle")
        assert page.title()
        assert page.locator(".brand").is_visible()
        assert page.locator("#panel-import").is_visible()
        page.screenshot(path=OUTPUT / "01-import.png", full_page=True)

        with page.expect_file_chooser() as chooser_info:
            page.locator("#chooseFilesBtn").click()
        chooser_info.value.set_files([fixture, second_fixture])
        expect(page.locator("#rawText")).to_have_value(re.compile("Second Resume Evidence File"))
        raw_text = page.locator("#rawText").input_value()
        assert "Resume Import Smoke Test" in raw_text
        assert "Second Resume Evidence File" in raw_text
        assert page.locator("#parseMeta").text_content().startswith("已提取 2/2")

        page.get_by_role("button", name="跳过，手动填写").click()
        page.wait_for_timeout(700)
        assert page.locator("#panel-profile").is_visible()
        page.locator('[data-profile-path="basics.name"]').fill("隔离测试用户")
        assert page.locator('[data-profile-path="basics.name"]').input_value() == "隔离测试用户"

        page.locator('[data-step="design"]').click()
        page.wait_for_timeout(900)
        assert page.locator("#panel-design").is_visible()
        assert page.locator("#resumePreview").is_visible()
        page.locator('[data-template="editorial"]').click()
        page.wait_for_timeout(500)
        page.screenshot(path=OUTPUT / "02-design.png", full_page=True)

        with page.expect_download() as docx_download:
            page.locator("#docxBtn").click()
        docx_download.value.save_as(OUTPUT / "ui-export.docx")

        with page.expect_download(timeout=120_000) as pdf_download:
            page.locator("#pdfBtn").click()
        pdf_download.value.save_as(OUTPUT / "ui-export.pdf")

        page.locator("#settingsBtn").click()
        assert page.locator("#settingsDialog").is_visible()
        page.screenshot(path=OUTPUT / "03-settings.png", full_page=True)
        browser.close()

    if console_errors:
        raise AssertionError("Browser errors: " + " | ".join(console_errors))
    print({"screenshots": 3, "docx": (OUTPUT / "ui-export.docx").stat().st_size, "pdf": (OUTPUT / "ui-export.pdf").stat().st_size})


if __name__ == "__main__":
    main()
