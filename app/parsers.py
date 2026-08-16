from __future__ import annotations

import io
import math
import zipfile
from collections import Counter
from pathlib import Path
from typing import Any
from xml.etree import ElementTree

import fitz
from docx import Document
from PIL import Image, ImageEnhance, ImageOps


_ocr_reader: Any = None
_rapid_ocr_reader: Any = None
MAX_ARCHIVE_ENTRIES = 2_000
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_ARCHIVE_MEMBER_BYTES = 30 * 1024 * 1024
MAX_IMAGE_PIXELS = 25_000_000
MAX_IMAGE_DIMENSION = 12_000
MAX_PDF_PAGES = 50
MAX_PDF_RENDER_PIXELS = 20_000_000


def _normalize(lines: list[str]) -> str:
    return "\n".join(line.strip() for line in lines if line and line.strip()).strip()


def _validate_office_archive(data: bytes) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ARCHIVE_ENTRIES:
                raise ValueError("Office 文件包含过多内部条目")
            total = 0
            for member in members:
                path = Path(member.filename.replace("\\", "/"))
                if member.flag_bits & 0x1:
                    raise ValueError("不支持加密的 Office 文件")
                if path.is_absolute() or ".." in path.parts:
                    raise ValueError("Office 文件包含不安全的内部路径")
                if member.file_size > MAX_ARCHIVE_MEMBER_BYTES:
                    raise ValueError("Office 文件中的单个内容过大")
                total += member.file_size
                if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise ValueError("Office 文件解压后的内容过大")
    except zipfile.BadZipFile as exc:
        raise ValueError("Office 文件结构损坏或格式不正确") from exc


def _open_safe_image(data: bytes) -> Image.Image:
    try:
        with Image.open(io.BytesIO(data)) as source:
            width, height = source.size
            if (
                width <= 0 or height <= 0
                or width > MAX_IMAGE_DIMENSION or height > MAX_IMAGE_DIMENSION
                or width * height > MAX_IMAGE_PIXELS
            ):
                raise ValueError("图片尺寸过大，请缩小后重试")
            source.load()
            return source.copy()
    except (Image.DecompressionBombError, OSError) as exc:
        raise ValueError("图片损坏或像素尺寸过大") from exc


def parse_docx(data: bytes) -> str:
    _validate_office_archive(data)
    document = Document(io.BytesIO(data))
    lines = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text.strip() for cell in row.cells))
    return _normalize(lines)


def _parse_xlsx(data: bytes) -> str:
    """Extract visible cell values from modern Excel files without executing formulas/macros."""
    _validate_office_archive(data)
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        shared: list[str] = []
        if "xl/sharedStrings.xml" in archive.namelist():
            root = ElementTree.fromstring(archive.read("xl/sharedStrings.xml"))
            shared = ["".join(node.itertext()) for node in root]
        lines: list[str] = []
        for name in sorted(item for item in archive.namelist() if item.startswith("xl/worksheets/sheet") and item.endswith(".xml"))[:20]:
            root = ElementTree.fromstring(archive.read(name))
            for row in root.iter():
                if not row.tag.endswith("}row"):
                    continue
                values: list[str] = []
                for cell in row:
                    if not cell.tag.endswith("}c"):
                        continue
                    value = next((node for node in cell if node.tag.endswith("}v")), None)
                    inline = next((node for node in cell if node.tag.endswith("}is")), None)
                    text = "".join(inline.itertext()) if inline is not None else (value.text if value is not None else "")
                    if cell.attrib.get("t") == "s" and text.isdigit() and int(text) < len(shared):
                        text = shared[int(text)]
                    values.append(text.strip())
                if any(values):
                    lines.append(" | ".join(values))
                if sum(len(line) for line in lines) > 80_000:
                    return _normalize(lines)[:80_000]
        return _normalize(lines)


def _dominant_accent(image: Image.Image) -> str:
    pixels = image.convert("RGB").resize((96, 96)).getdata()
    candidates = []
    for red, green, blue in pixels:
        spread = max(red, green, blue) - min(red, green, blue)
        brightness = (red + green + blue) / 3
        if spread > 28 and 35 < brightness < 220:
            candidates.append((red // 24 * 24, green // 24 * 24, blue // 24 * 24))
    color = Counter(candidates).most_common(1)[0][0] if candidates else (31, 58, 52)
    return "#%02x%02x%02x" % color


def analyze_template(filename: str, data: bytes) -> dict[str, Any]:
    """Convert an imported reference into the app's constrained, editable style model."""
    extension = Path(filename).suffix.lower()
    if len(data) > 20 * 1024 * 1024:
        raise ValueError("模板文件不能超过 20 MB")
    spec: dict[str, Any] = {
        "name": Path(filename).stem[:80] or "自定义模板",
        "source_type": extension.lstrip("."),
        "accent": "#1f3a34",
        "font_family": "Microsoft YaHei",
        "layout": "single",
        "header_align": "left",
        "density": "compact",
        "has_photo": False,
        "notes": [],
    }
    if extension == ".docx":
        _validate_office_archive(data)
        document = Document(io.BytesIO(data))
        fonts = Counter()
        colors = Counter()
        for paragraph in document.paragraphs[:120]:
            for run in paragraph.runs:
                if run.font.name:
                    fonts[run.font.name] += max(1, len(run.text))
                if run.font.color and run.font.color.rgb:
                    colors[str(run.font.color.rgb)] += max(1, len(run.text))
        if fonts:
            spec["font_family"] = fonts.most_common(1)[0][0]
        if colors:
            spec["accent"] = "#" + colors.most_common(1)[0][0].lower()
        spec["layout"] = "sidebar" if any(len(table.columns) >= 2 for table in document.tables) else "single"
        spec["has_photo"] = bool(document.inline_shapes)
        if document.paragraphs and document.paragraphs[0].alignment == 1:
            spec["header_align"] = "center"
    elif extension == ".pdf":
        document = fitz.open(stream=data, filetype="pdf")
        if not len(document):
            raise ValueError("PDF 模板没有可读取的页面")
        page = document[0]
        pixmap = page.get_pixmap(matrix=fitz.Matrix(1, 1), alpha=False)
        spec["accent"] = _dominant_accent(Image.open(io.BytesIO(pixmap.tobytes("png"))))
        blocks = [block for block in page.get_text("blocks") if str(block[4]).strip()]
        left = sum(1 for block in blocks if block[0] < page.rect.width * .38)
        right = sum(1 for block in blocks if block[0] > page.rect.width * .42)
        spec["layout"] = "sidebar" if left >= 3 and right >= 3 else "single"
        spec["has_photo"] = bool(page.get_images(full=True))
    elif extension in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
        image = _open_safe_image(data)
        spec["accent"] = _dominant_accent(image)
        spec["has_photo"] = True
        spec["notes"].append("图片模板已提取主色与版式倾向；文字区域将按可编辑模块重建。")
    else:
        raise ValueError("模板支持 DOCX、PDF、PNG、JPG、WEBP 和 BMP")
    spec["notes"].append("已重建为可编辑模板；不会导入模板中的示例人物与简历文字。")
    return {"template_spec": spec}


def _get_ocr_reader():
    global _ocr_reader
    if _ocr_reader is None:
        try:
            import easyocr

            _ocr_reader = easyocr.Reader(["ch_sim", "en"], gpu=False, verbose=False)
        except Exception as exc:
            raise RuntimeError(
                "EasyOCR 初始化失败。请确认本机模型已安装且可用。"
            ) from exc
    return _ocr_reader


def _get_rapid_ocr_reader():
    global _rapid_ocr_reader
    if _rapid_ocr_reader is None:
        from rapidocr import RapidOCR

        _rapid_ocr_reader = RapidOCR()
    return _rapid_ocr_reader


def _enhance_for_ocr(image: Image.Image) -> Image.Image:
    rgb = image.convert("RGB")
    if rgb.width < 1800:
        scale = min(2.0, 1800 / max(rgb.width, 1))
        rgb = rgb.resize((round(rgb.width * scale), round(rgb.height * scale)), Image.Resampling.LANCZOS)
    gray = ImageOps.grayscale(rgb)
    gray = ImageOps.autocontrast(gray, cutoff=1)
    gray = ImageEnhance.Sharpness(gray).enhance(1.35)
    return gray.convert("RGB")


def _rapid_ocr(image: Image.Image) -> tuple[str, float]:
    import numpy as np

    output = _get_rapid_ocr_reader()(np.array(image), text_score=0.45)
    lines = list(output.txts or [])
    scores = list(output.scores or [])
    text = _normalize([str(line) for line in lines])
    confidence = sum(float(score) for score in scores) / len(scores) if scores else 0.0
    return text, confidence


def _ocr_image_with_confidence(image: Image.Image) -> tuple[str, float | None]:
    import numpy as np

    try:
        original = image.convert("RGB")
        text, confidence = _rapid_ocr(original)
        if confidence < 0.88 or len(text) < 80:
            enhanced_text, enhanced_confidence = _rapid_ocr(_enhance_for_ocr(original))
            original_quality = confidence + min(len(text) / 2000, 0.12)
            enhanced_quality = enhanced_confidence + min(len(enhanced_text) / 2000, 0.12)
            if enhanced_quality > original_quality:
                text, confidence = enhanced_text, enhanced_confidence
        if text:
            return text, confidence
    except Exception:
        pass

    reader = _get_ocr_reader()
    result = reader.readtext(
        np.array(_enhance_for_ocr(image)),
        detail=1,
        paragraph=False,
        decoder="beamsearch",
    )
    lines = [str(item[1]) for item in result if len(item) >= 3]
    scores = [float(item[2]) for item in result if len(item) >= 3]
    return _normalize(lines), (sum(scores) / len(scores) if scores else None)


def _ocr_image(image: Image.Image) -> str:
    return _ocr_image_with_confidence(image)[0]


def parse_pdf(data: bytes) -> tuple[str, bool]:
    text, used_ocr, _ = parse_pdf_details(data)
    return text, used_ocr


def parse_pdf_details(data: bytes) -> tuple[str, bool, float | None]:
    document = fitz.open(stream=data, filetype="pdf")
    if len(document) > MAX_PDF_PAGES:
        raise ValueError(f"PDF 最多支持 {MAX_PDF_PAGES} 页")
    page_text = [page.get_text("text").strip() for page in document]
    text = _normalize(page_text)
    meaningful_chars = sum(ch.isalnum() for ch in text)
    if meaningful_chars >= max(80, len(document) * 30):
        return text, False, None

    ocr_pages: list[str] = []
    confidences: list[float] = []
    for page in document:
        target_scale = 2.0
        estimated_pixels = max(page.rect.width, 1) * max(page.rect.height, 1) * target_scale * target_scale
        if estimated_pixels > MAX_PDF_RENDER_PIXELS:
            target_scale *= math.sqrt(MAX_PDF_RENDER_PIXELS / estimated_pixels)
        pixmap = page.get_pixmap(matrix=fitz.Matrix(target_scale, target_scale), alpha=False)
        image = Image.open(io.BytesIO(pixmap.tobytes("png")))
        page_text, confidence = _ocr_image_with_confidence(image)
        ocr_pages.append(page_text)
        if confidence is not None:
            confidences.append(confidence)
    return _normalize(ocr_pages), True, (sum(confidences) / len(confidences) if confidences else None)


def parse_image(data: bytes) -> str:
    return _ocr_image(_open_safe_image(data))


def parse_file(filename: str, data: bytes) -> dict[str, Any]:
    extension = Path(filename).suffix.lower()
    if len(data) > 20 * 1024 * 1024:
        raise ValueError("文件不能超过 20 MB")
    if extension == ".docx":
        text = parse_docx(data)
        used_ocr = False
        confidence = None
    elif extension == ".pdf":
        text, used_ocr, confidence = parse_pdf_details(data)
    elif extension in {".png", ".jpg", ".jpeg", ".webp", ".bmp"}:
        text, confidence = _ocr_image_with_confidence(_open_safe_image(data))
        used_ocr = True
    elif extension in {".txt", ".md", ".csv", ".json", ".py", ".js", ".ts", ".tsx", ".jsx", ".html", ".css", ".sql", ".yaml", ".yml"}:
        text = data.decode("utf-8", errors="replace")[:100_000]
        used_ocr = False
        confidence = None
    elif extension == ".xlsx":
        text = _parse_xlsx(data)
        used_ocr = False
        confidence = None
    else:
        raise ValueError("支持 DOCX、PDF、图片、XLSX、CSV、Markdown 和常见代码文本文件")
    if not text.strip():
        raise ValueError("未能从文件中识别出文字")
    return {"filename": filename, "text": text, "used_ocr": used_ocr, "ocr_confidence": confidence}
