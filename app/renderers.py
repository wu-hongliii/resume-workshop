from __future__ import annotations

import base64
import html
import io
from typing import Any, Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from . import runtime

RESUME_CSS = runtime.RESOURCE_ROOT / "web" / "resume.css"
SECTION_LABELS = {
    "summary": "个人概述",
    "experience": "工作经历",
    "projects": "项目经历",
    "education": "教育背景",
    "skills": "专业技能",
    "certificates": "证书与荣誉",
    "languages": "语言能力",
}
PROTECTED_BASIC_FIELDS = ("name", "phone", "email", "city", "links", "photo")
MIN_READABLE_ONE_PAGE_SCALE = 0.75


class ResumeLayoutError(ValueError):
    pass


def effective_profile(payload: dict[str, Any]) -> dict[str, Any]:
    base = payload.get("profile", payload)
    resume = payload.get("resume", {})
    tailored = resume.get("tailored_profile")
    if not isinstance(tailored, dict):
        return base
    import copy

    result = copy.deepcopy(tailored)
    result.setdefault("basics", {})
    base_basics = base.get("basics", {})
    for field in PROTECTED_BASIC_FIELDS:
        result["basics"][field] = base_basics.get(field, "")
    return result


def _e(value: Any) -> str:
    return html.escape(str(value or ""))


def _items(value: Any) -> list[str]:
    if isinstance(value, list):
        return [str(item).strip() for item in value if str(item).strip()]
    if isinstance(value, str):
        return [line.strip(" •-\t") for line in value.splitlines() if line.strip(" •-\t")]
    return []


def _inline(value: Any) -> str:
    if isinstance(value, list):
        return "、".join(str(item).strip() for item in value if str(item).strip())
    return str(value or "").strip()


def _bullets(items: Iterable[Any]) -> str:
    values = [f"<li>{_e(item)}</li>" for item in items if str(item).strip()]
    return f"<ul>{''.join(values)}</ul>" if values else ""


def _date_range(item: dict[str, Any]) -> str:
    start, end = item.get("start", ""), item.get("end", "")
    return " — ".join(part for part in (str(start).strip(), str(end).strip()) if part)


def _render_section(key: str, profile: dict[str, Any]) -> str:
    title = SECTION_LABELS.get(key, key)
    body = ""
    if key == "summary":
        body = f'<p class="summary">{_e(profile.get("basics", {}).get("summary", ""))}</p>'
    elif key in {"experience", "projects", "education"}:
        cards = []
        for item in profile.get(key, []):
            if key == "experience":
                heading = item.get("role") or item.get("company")
                sub = item.get("company") if item.get("role") else ""
                meta = " · ".join(filter(None, [_inline(sub), _inline(item.get("location", ""))]))
                detail = _bullets(_items(item.get("bullets", [])))
            elif key == "projects":
                heading = item.get("name") or item.get("role")
                meta = " · ".join(filter(None, [_inline(item.get("role", "")), _inline(item.get("technologies", ""))]))
                detail = _bullets(_items(item.get("bullets", [])))
            else:
                heading = item.get("school") or item.get("degree")
                meta = " · ".join(filter(None, [_inline(item.get("degree", "")), _inline(item.get("major", ""))]))
                detail = ""
            cards.append(
                '<article class="entry">'
                f'<div class="entry-head"><div><h3>{_e(heading)}</h3><p>{_e(meta)}</p></div>'
                f'<time>{_e(_date_range(item))}</time></div>{detail}</article>'
            )
        body = "".join(cards)
    elif key == "skills":
        rows = []
        for item in profile.get("skills", []):
            values = item.get("items", [])
            if isinstance(values, list):
                values = "、".join(str(value) for value in values)
            rows.append(
                f'<div class="skill-row"><strong>{_e(item.get("category", "技能"))}</strong>'
                f'<span>{_e(values)}</span></div>'
            )
        body = '<div class="skill-list">' + "".join(rows) + "</div>"
    else:
        body = _bullets(_items(profile.get(key, [])))
    if not body or body in {'<p class="summary"></p>', '<div class="skill-list"></div>'}:
        return ""
    return f'<section class="resume-section section-{_e(key)}"><h2>{_e(title)}</h2>{body}</section>'


def render_resume_html(payload: dict[str, Any]) -> str:
    profile = effective_profile(payload)
    resume = payload.get("resume", {})
    basics = profile.get("basics", {})
    order = resume.get("section_order") or list(SECTION_LABELS)
    sections = "".join(_render_section(key, profile) for key in order)
    contact = " · ".join(
        str(value).strip()
        for value in [basics.get("phone"), basics.get("email"), basics.get("city"), basics.get("links")]
        if str(value or "").strip()
    )
    photo = basics.get("photo", "")
    photo_html = f'<img class="portrait" src="{_e(photo)}" alt="个人照片">' if photo else ""
    css = RESUME_CSS.read_text(encoding="utf-8")
    template = resume.get("template", "ats")
    custom = resume.get("custom_template") if isinstance(resume.get("custom_template"), dict) else {}
    accent = str(custom.get("accent", "#1f3a34")) if template == "custom" else ""
    font = str(custom.get("font_family", "Microsoft YaHei")) if template == "custom" else ""
    layout = str(custom.get("layout", "single")) if template == "custom" else "single"
    header_align = str(custom.get("header_align", "left")) if template == "custom" else "left"
    custom_style = f' style="--accent:{_e(accent)};--custom-font:{_e(font)}"' if template == "custom" else ""
    page_mode = resume.get("page_mode", "one")
    return f"""<!doctype html>
<html lang="zh-CN"><head><meta charset="utf-8"><style>{css}</style></head>
<body><main class="resume-page template-{_e(template)} layout-{_e(layout)} header-{_e(header_align)} page-{_e(page_mode)}"{custom_style}><div class="resume-fit">
<header class="resume-header">{photo_html}<div class="identity"><h1>{_e(basics.get('name') or '姓名')}</h1>
<p class="target-role">{_e(basics.get('target_role'))}</p><p class="contact">{_e(contact)}</p></div></header>
<div class="resume-content">{sections}</div></div></main>
<script>
window.fitResumeToOnePage = function() {{
  const page = document.querySelector('.resume-page');
  const fit = document.querySelector('.resume-fit');
  fit.style.transform = 'none';
  fit.style.width = '100%';
  page.dataset.fitScale = '1';
  if (!page.classList.contains('page-one')) return 1;
  const style = getComputedStyle(page);
  const available = page.clientHeight - parseFloat(style.paddingTop) - parseFloat(style.paddingBottom);
  const scale = Math.min(1, available / Math.max(1, fit.scrollHeight));
  fit.style.transform = `scale(${{scale}})`;
  page.dataset.fitScale = scale.toFixed(3);
  return scale;
}};
const runFit = () => requestAnimationFrame(() => window.fitResumeToOnePage());
window.addEventListener('load', runFit);
window.addEventListener('resize', runFit);
if (document.fonts && document.fonts.ready) document.fonts.ready.then(runFit);
</script></body></html>"""


async def render_pdf(payload: dict[str, Any]) -> bytes:
    from playwright.async_api import async_playwright

    markup = render_resume_html(payload)
    async with async_playwright() as playwright:
        browser = await playwright.chromium.launch(headless=True)
        try:
            page = await browser.new_page()
            await page.set_content(markup, wait_until="load")
            scale = await page.evaluate("window.fitResumeToOnePage()")
            if payload.get("resume", {}).get("page_mode", "one") == "one" and scale < MIN_READABLE_ONE_PAGE_SCALE:
                raise ResumeLayoutError(
                    f"内容缩放后只有 {round(scale * 100)}%，低于可读下限。请精简内容或选择允许多页。"
                )
            await page.emulate_media(media="print")
            return await page.pdf(
                format="A4",
                print_background=True,
                prefer_css_page_size=True,
                margin={"top": "0", "right": "0", "bottom": "0", "left": "0"},
            )
        finally:
            await browser.close()


def _set_run_font(run, name: str = "Microsoft YaHei", size: float = 9.5, bold: bool = False, color: str = "252722"):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def _paragraph_border(paragraph, color: str = "B88B4A", size: int = 8) -> None:
    p_pr = paragraph._p.get_or_add_pPr()
    borders = p_pr.find(qn("w:pBdr"))
    if borders is None:
        borders = OxmlElement("w:pBdr")
        p_pr.append(borders)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), color)
    borders.append(bottom)


def _add_heading(document: Document, text: str, accent: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(4)
    _set_run_font(paragraph.add_run(text), size=11, bold=True, color=accent)
    _paragraph_border(paragraph, accent, 6)


def _add_bullets(document: Document, bullets: Any) -> None:
    for bullet in _items(bullets):
        paragraph = document.add_paragraph(style="List Bullet")
        paragraph.paragraph_format.left_indent = Cm(0.5)
        paragraph.paragraph_format.first_line_indent = Cm(-0.25)
        paragraph.paragraph_format.space_after = Pt(2)
        _set_run_font(paragraph.add_run(bullet), size=9.3)


def _add_entry(document: Document, key: str, item: dict[str, Any], accent: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(3)
    paragraph.paragraph_format.space_after = Pt(1)
    if key == "experience":
        heading = item.get("role") or item.get("company", "")
        meta = " · ".join(filter(None, [item.get("company", ""), item.get("location", "")]))
    elif key == "projects":
        heading = item.get("name") or item.get("role", "")
        meta = " · ".join(filter(None, [item.get("role", ""), item.get("technologies", "")]))
    else:
        heading = item.get("school") or item.get("degree", "")
        meta = " · ".join(filter(None, [item.get("degree", ""), item.get("major", "")]))
    _set_run_font(paragraph.add_run(str(heading)), size=10, bold=True, color=accent)
    date = _date_range(item)
    if date:
        _set_run_font(paragraph.add_run(f"    {date}"), size=8.5, color="60635C")
    if meta:
        meta_paragraph = document.add_paragraph()
        meta_paragraph.paragraph_format.space_after = Pt(1)
        _set_run_font(meta_paragraph.add_run(str(meta)), size=8.8, color="60635C")
    if key != "education":
        _add_bullets(document, item.get("bullets", []))


def render_docx(payload: dict[str, Any]) -> bytes:
    profile = effective_profile(payload)
    resume = payload.get("resume", {})
    basics = profile.get("basics", {})
    template = resume.get("template", "ats")
    accents = {"ats": "1F3A34", "editorial": "8B4A32", "creative": "234B78"}
    custom = resume.get("custom_template") if isinstance(resume.get("custom_template"), dict) else {}
    accent = str(custom.get("accent", "#1f3a34")).lstrip("#").upper() if template == "custom" else accents.get(template, accents["ats"])
    if len(accent) != 6 or any(character not in "0123456789ABCDEF" for character in accent):
        accent = accents["ats"]

    document = Document()
    section = document.sections[0]
    section.start_type = WD_SECTION_START.NEW_PAGE
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(1.35)
    section.bottom_margin = Cm(1.35)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)

    normal = document.styles["Normal"]
    normal.paragraph_format.space_after = Pt(0)
    normal.font.name = "Microsoft YaHei"
    normal.font.size = Pt(9.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    name = document.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER if template == "ats" or (template == "custom" and custom.get("header_align") == "center") else WD_ALIGN_PARAGRAPH.LEFT
    name.paragraph_format.space_after = Pt(2)
    _set_run_font(name.add_run(str(basics.get("name") or "姓名")), size=22, bold=True, color=accent)
    role = document.add_paragraph()
    role.alignment = name.alignment
    role.paragraph_format.space_after = Pt(3)
    _set_run_font(role.add_run(str(basics.get("target_role", ""))), size=10.5, color="60635C")
    contact_values = [basics.get("phone"), basics.get("email"), basics.get("city"), basics.get("links")]
    contact = document.add_paragraph()
    contact.alignment = name.alignment
    contact.paragraph_format.space_after = Pt(5)
    _set_run_font(contact.add_run(" · ".join(str(v) for v in contact_values if v)), size=8.5, color="60635C")

    order = resume.get("section_order") or list(SECTION_LABELS)
    for key in order:
        if key == "summary":
            value = basics.get("summary", "")
            if value:
                _add_heading(document, SECTION_LABELS[key], accent)
                paragraph = document.add_paragraph()
                paragraph.paragraph_format.line_spacing = 1.12
                _set_run_font(paragraph.add_run(str(value)))
        elif key in {"experience", "projects", "education"}:
            values = profile.get(key, [])
            if values:
                _add_heading(document, SECTION_LABELS[key], accent)
                for item in values:
                    _add_entry(document, key, item, accent)
        elif key == "skills":
            values = profile.get("skills", [])
            if values:
                _add_heading(document, SECTION_LABELS[key], accent)
                for item in values:
                    values_text = item.get("items", [])
                    if isinstance(values_text, list):
                        values_text = "、".join(str(value) for value in values_text)
                    paragraph = document.add_paragraph()
                    _set_run_font(paragraph.add_run(f"{item.get('category', '技能')}："), bold=True, color=accent)
                    _set_run_font(paragraph.add_run(str(values_text)))
        else:
            values = profile.get(key, [])
            if values:
                _add_heading(document, SECTION_LABELS.get(key, key), accent)
                _add_bullets(document, values)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def safe_filename(name: str, extension: str) -> str:
    cleaned = "".join(char for char in name if char not in '<>:"/\\|?*').strip() or "简历"
    return f"{cleaned[:60]}.{extension}"
